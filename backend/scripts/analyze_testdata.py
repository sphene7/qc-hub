"""Analyze testdata Excel files vs system capabilities."""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.parser import parse_workbook
from app.semantics import build_workbook_spec
import openpyxl

TESTDATA = Path(__file__).resolve().parents[2] / "testdata"


def sheet_overview_wb(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    info = {"file": path.name, "sheets": []}
    for ws in wb.worksheets:
        info["sheets"].append({
            "name": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "state": getattr(ws, "sheet_state", "visible"),
        })
    wb.close()
    return info


def main() -> None:
    results = []
    for f in sorted(TESTDATA.glob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        entry = {"file": f.name, "raw": sheet_overview_wb(f)}
        try:
            parsed = parse_workbook(f, f.name)
            spec = build_workbook_spec(parsed)
            sheets = []
            for s in spec["sheets"]:
                sheets.append({
                    "name": s["name"],
                    "kind": s["kind"],
                    "row_count": s["row_count"],
                    "col_count": s["col_count"],
                    "primary": s.get("primary"),
                    "empty": s.get("empty"),
                    "formula_count": s["formula"].get("count", 0),
                    "formula_functions": s["formula"].get("functions", []),
                    "cross_sheets": s["formula"].get("cross_sheets", []),
                    "columns": [c["key"] for c in s.get("columns", [])],
                    "column_roles": {c["key"]: c["role"] for c in s.get("columns", [])},
                    "kpis": [k["label"] for k in s.get("kpis", [])],
                    "charts": [c["title"] for c in s.get("charts", [])],
                    "filters": [f["label"] for f in s.get("filters", [])],
                    "joins": s.get("joins", []),
                    "note": s.get("note"),
                })
            entry["parsed"] = {
                "title": spec.get("title"),
                "primary_sheet": spec.get("primary_sheet"),
                "sheets": sheets,
                "smart_view": spec.get("smart_view", {}),
            }
        except Exception as e:
            entry["error"] = str(e)
        results.append(entry)

    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
