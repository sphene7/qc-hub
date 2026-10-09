"""Load Excel workbooks into typed tables, detecting headers, dates, and used range."""

from __future__ import annotations

import math
import re
import warnings
from datetime import datetime, timedelta, date
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.utils.datetime import from_excel

from .formulas import collect_sheet_formulas, index_to_col_letters

MAX_COLS = 250
MAX_ROWS = 20000
EMPTY_STOP = 25
HEADER_SCAN = 25

EXCEL_EPOCH = datetime(1899, 12, 30)


def parse_workbook(path: str | Path, filename: str | None = None) -> dict[str, Any]:
    path = Path(path)
    filename = filename or path.name
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb_f = openpyxl.load_workbook(path, data_only=False, read_only=True)
        wb_v = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = []
        for ws_f in wb_f.worksheets:
            name = ws_f.title
            hidden = getattr(ws_f, "sheet_state", "visible") != "visible"
            ws_v = wb_v[name] if name in wb_v.sheetnames else None
            sheet = _parse_sheet(ws_f, ws_v, name, hidden)
            sheets.append(sheet)
            # read_only worksheets can only be iterated once; already consumed here
    finally:
        wb_f.close()
        wb_v.close()

    derived = []
    for s in sheets:
        extra = _maybe_unpivot(s)
        if extra:
            derived.append(extra)
    sheets.extend(derived)

    return {
        "filename": filename,
        "title": _pretty_title(filename),
        "sheets": sheets,
    }


def _parse_sheet(ws_f, ws_v, name: str, hidden: bool) -> dict[str, Any]:
    raw_rows, formulas, last_row, used_cols = _read_rows(ws_f, ws_v)
    header_row = _detect_header_row(raw_rows)
    headers = _headers_from(raw_rows, header_row, used_cols)
    data_start = header_row  # 1-based header; data is next
    records: list[dict[str, Any]] = []
    for abs_i, row in enumerate(raw_rows[header_row:], start=header_row + 1):
        if _row_empty(row):
            continue
        rec = {}
        for i, h in enumerate(headers):
            rec[h] = row[i] if i < len(row) else None
        rec["__row"] = abs_i
        records.append(rec)

    columns = _profile_columns(headers, records, formulas, header_row)
    df = _to_dataframe(columns, records)
    formula_info = collect_sheet_formulas(
        formulas,
        name,
        header_row,
        last_row,
        used_cols,
        headers,
    )
    kind = _classify_sheet(name, hidden, df, formula_info, columns)

    preview = _json_records(df.head(8))
    return {
        "name": name,
        "hidden": hidden,
        "kind": kind,
        "header_row": header_row,
        "row_count": int(len(df)),
        "col_count": len(headers),
        "columns": columns,
        "formula": formula_info,
        "preview": preview,
        "df": df,
        "empty": len(df) == 0,
        "note": _sheet_note(name, kind, len(df), formula_info),
    }


def _read_rows(ws_f, ws_v):
    """One-pass read (read_only worksheets cannot be iterated twice).

    Returns value-rows (0-indexed), formula map (1-based), last row, used col count.
    """
    declared = ws_f.max_column or 1
    cap = min(max(declared, 1), MAX_COLS)
    formula_rows: dict[int, list] = {}
    formulas: dict[tuple[int, int], str] = {}
    empty = 0
    last_data = 1
    last_col = 1
    for i, row in enumerate(
        ws_f.iter_rows(min_row=1, max_row=MAX_ROWS, max_col=cap, values_only=False),
        start=1,
    ):
        vals = []
        has_anchor = False
        for c_i, cell in enumerate(row, start=1):
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                formulas[(i, c_i)] = v
                vals.append(None)
                last_col = max(last_col, c_i)
            else:
                vals.append(v)
                if v not in (None, ""):
                    has_anchor = True
                    last_col = max(last_col, c_i)
        formula_rows[i] = vals
        if has_anchor:
            empty = 0
            last_data = i
        else:
            empty += 1
            if empty >= EMPTY_STOP and i > 8:
                break

    used_cols = min(max(last_col, 1), cap)
    value_rows: dict[int, list] = {}
    if ws_v is not None:
        for i, row in enumerate(
            ws_v.iter_rows(min_row=1, max_row=last_data, max_col=used_cols, values_only=True),
            start=1,
        ):
            value_rows[i] = list(row)

    merged: list[list[Any]] = []
    for r in range(1, last_data + 1):
        raw = value_rows.get(r) or formula_rows.get(r) or [None] * used_cols
        if len(raw) < used_cols:
            raw = list(raw) + [None] * (used_cols - len(raw))
        merged.append(raw[:used_cols])
    return merged, formulas, last_data, used_cols


def _detect_header_row(rows: list[list[Any]]) -> int:
    """Return 1-based header row index."""
    best_i, best_score = 1, -999
    limit = min(len(rows), HEADER_SCAN)
    for i in range(limit):
        row = rows[i]
        nonempty = [v for v in row if v not in (None, "")]
        if not nonempty:
            continue
        strings = [v for v in nonempty if isinstance(v, str)]
        nums = [v for v in nonempty if isinstance(v, (int, float)) and not isinstance(v, bool)]
        score = 0
        score += len(strings) * 3
        score += len({str(s).strip() for s in strings})
        score -= len(nums) * 2
        if len(nonempty) == 1 and isinstance(nonempty[0], str) and len(nonempty[0]) > 12:
            score -= 20  # title row
        # field-name bonus
        for s in strings:
            t = str(s)
            if any(k in t for k in ("日期", "月份", "数量", "机种", "部门", "状态", "判定", "率", "序号")):
                score += 4
            if len(t) > 40:
                score -= 6
        # next row looks like data
        if i + 1 < len(rows):
            nxt = [v for v in rows[i + 1] if v not in (None, "")]
            nxt_num = sum(1 for v in nxt if isinstance(v, (int, float)) and not isinstance(v, bool))
            if nxt_num >= 2:
                score += 6
        if score > best_score:
            best_score, best_i = score, i + 1
    return best_i


def _headers_from(rows: list[list[Any]], header_row: int, used_cols: int) -> list[str]:
    row = rows[header_row - 1] if header_row - 1 < len(rows) else []
    names = []
    seen: dict[str, int] = {}
    for i in range(used_cols):
        v = row[i] if i < len(row) else None
        name = _clean_header(v) or f"列{index_to_col_letters(i + 1)}"
        n = seen.get(name, 0) + 1
        seen[name] = n
        names.append(name if n == 1 else f"{name}_{n}")
    # drop trailing auto-generated empty columns if original was empty
    while len(names) > 1 and names[-1].startswith("列") and _is_empty_col(rows, header_row, len(names) - 1):
        names.pop()
    return names


def _is_empty_col(rows: list[list[Any]], header_row: int, idx: int) -> bool:
    for row in rows[header_row:]:
        if idx < len(row) and row[idx] not in (None, ""):
            return False
    return True


def _clean_header(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).replace("\n", "").replace("\r", "")
    s = re.sub(r"\s+", "", s).strip()
    return s


def _row_empty(row: list[Any]) -> bool:
    return all(v in (None, "") for v in row)


def _excel_serial_to_date(v: float) -> date | None:
    try:
        fv = float(v)
    except (TypeError, ValueError):
        return None
    if not (20000 <= fv <= 80000) or abs(fv - round(fv)) > 1e-6:
        # allow time fractions
        if not (20000 <= fv <= 80000):
            return None
    try:
        dt = EXCEL_EPOCH + timedelta(days=fv)
        return dt.date()
    except OverflowError:
        return None


def _looks_like_date_header(name: str) -> bool:
    n = (name or "").strip()
    if n in ("抽检状况",):
        return True
    return any(k in name for k in ("日期", "时间", "日")) and "时间段" not in name and "天数" not in name


def _profile_columns(
    headers: list[str],
    records: list[dict[str, Any]],
    formulas: dict[tuple[int, int], str],
    header_row: int,
) -> list[dict[str, Any]]:
    columns = []
    n = len(records)
    for i, h in enumerate(headers):
        vals = [r.get(h) for r in records]
        nonempty = [v for v in vals if v not in (None, "")]
        fill = len(nonempty) / n if n else 0
        date_hits = 0
        converted: list[Any] = []
        for v in vals:
            if v in (None, ""):
                converted.append(None)
                continue
            if isinstance(v, datetime):
                converted.append(v.date().isoformat())
                date_hits += 1
                continue
            if isinstance(v, date):
                converted.append(v.isoformat())
                date_hits += 1
                continue
            if (h or "").strip() == "抽检状况" and isinstance(v, (int, float)):
                d = _excel_serial_to_date(v)
                if d:
                    converted.append(f"{d.month}月{d.day}日")
                    date_hits += 1
                    continue
            if _looks_like_date_header(h) and isinstance(v, (int, float)):
                d = _excel_serial_to_date(v)
                if d:
                    converted.append(d.isoformat())
                    date_hits += 1
                    continue
            if isinstance(v, float) and math.isnan(v):
                converted.append(None)
                continue
            if isinstance(v, str):
                s = v.strip()
                if s in ("#N/A", "#DIV/0!", "#VALUE!", "#REF!", "-", "－", "/"):
                    converted.append(None)
                    continue
                converted.append(s)
                continue
            converted.append(v)

        for r, nv in zip(records, converted):
            r[h] = nv

        nn = [v for v in converted if v not in (None, "")]
        uniq = list(dict.fromkeys(nn))
        types = _type_counts(nn)
        inferred = _infer_type(h, types, date_hits, len(nn), n)
        cardinality = len(uniq)
        formula_col = sum(
            1
            for (r, c) in formulas
            if c == i + 1 and r > header_row
        )
        columns.append(
            {
                "key": h,
                "label": h,
                "index": i + 1,
                "type": inferred,
                "fill": round(fill, 3),
                "unique": cardinality,
                "samples": [ _sample(v) for v in uniq[:8] ],
                "formula_cells": formula_col,
            }
        )
    return columns


def _type_counts(values: list[Any]) -> dict[str, int]:
    counts = {"number": 0, "string": 0, "date": 0, "bool": 0}
    for v in values:
        if isinstance(v, bool):
            counts["bool"] += 1
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            counts["number"] += 1
        elif isinstance(v, str) and re.match(r"^\d{4}-\d{2}-\d{2}", v):
            counts["date"] += 1
        else:
            counts["string"] += 1
    return counts


def _infer_type(name: str, types: dict[str, int], date_hits: int, filled: int, n: int) -> str:
    if filled == 0:
        return "empty"
    if date_hits >= max(filled * 0.6, 1) or types["date"] >= filled * 0.6:
        return "date"
    if types["number"] >= filled * 0.7:
        if any(k in name for k in ("率", "比率", "占比")):
            return "percent"
        return "number"
    # month like 8 / 1月
    if "月份" in name or name.endswith("月"):
        return "category"
    if types["string"] >= filled * 0.5:
        return "string"
    return "string"


def _sample(v: Any) -> Any:
    if isinstance(v, float):
        return round(v, 6)
    if isinstance(v, str):
        return v[:40]
    return v


def _to_dataframe(columns: list[dict[str, Any]], records: list[dict[str, Any]]) -> pd.DataFrame:
    keys = [c["key"] for c in columns]
    if not records:
        return pd.DataFrame(columns=keys)
    data = [{k: r.get(k) for k in keys} for r in records]
    df = pd.DataFrame(data, columns=keys)
    for col in columns:
        if col["type"] in ("number", "percent"):
            df[col["key"]] = pd.to_numeric(df[col["key"]], errors="coerce")
        elif col["type"] == "date":
            df[col["key"]] = pd.to_datetime(df[col["key"]], errors="coerce")
    return df


def _classify_sheet(name: str, hidden: bool, df: pd.DataFrame, formula_info: dict, columns: list) -> str:
    n = name.strip()
    if hidden:
        return "hidden"
    if any(k in n for k in ("填写说明", "说明")) and len(df) <= 5:
        return "instruction"
    if any(k in n for k in ("代码", "定义", "对照")):
        return "lookup"
    funcs = {f["name"] for f in formula_info.get("functions") or []}
    formula_density = formula_info.get("count") or 0
    if len(df) <= 40 and formula_density >= 8 and funcs & {
        "COUNTIF", "COUNTIFS", "SUMIF", "SUMIFS", "SUMPRODUCT",
    }:
        return "dashboard"
    if len(df) == 0:
        return "empty"
    filled_cols = sum(1 for c in columns if c["fill"] > 0.05)
    if filled_cols >= 4 and len(df) >= 5:
        return "fact"
    if filled_cols >= 3:
        return "table"
    return "other"


def _sheet_note(name: str, kind: str, rows: int, formula_info: dict) -> str | None:
    if kind == "empty" and (formula_info.get("count") or 0) == 0:
        return "此表只有表头、没有明细数据（常见于 WPS 未把值写入单元格）。已保留列结构和其它表上的公式关系。"
    if kind == "dashboard":
        return "这是汇总看板表，数字多半由 COUNTIF / SUMIF 从明细表算出来。"
    if kind == "lookup":
        return "这是对照/代码表，明细表里的 VLOOKUP 会引用它。"
    if kind == "instruction":
        return "填写说明，不参与统计。"
    return None


def _json_records(df: pd.DataFrame) -> list[dict[str, Any]]:
    if df is None or df.empty:
        return []
    out = []
    for rec in df.to_dict(orient="records"):
        item = {}
        for k, v in rec.items():
            if v is None or (isinstance(v, float) and math.isnan(v)) or pd.isna(v):
                item[k] = None
            elif hasattr(v, "isoformat"):
                item[k] = v.isoformat()[:19]
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                item[k] = float(v) if isinstance(v, float) else int(v)
            else:
                item[k] = v if not hasattr(v, "item") else v.item()
        out.append(item)
    return out


def _pretty_title(filename: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"\d{8,}", "", stem)
    stem = re.sub(r"NEW", "", stem, flags=re.I)
    return stem.strip("-_ ") or filename


UNPIVOT_SUFFIXES = ("抽检数", "不良数", "良品数", "不良品")


def _maybe_unpivot(sheet: dict[str, Any]) -> dict[str, Any] | None:
    if sheet.get("kind") not in ("fact", "table"):
        return None
    df: pd.DataFrame = sheet["df"]
    if df.empty or len(df.columns) < 8:
        return None
    groups: dict[str, dict[str, str]] = {}
    used = set()
    for col in df.columns:
        for suf in UNPIVOT_SUFFIXES:
            if col.endswith(suf):
                prefix = re.sub(r"[-_—\s]+$", "", col[: -len(suf)])
                if not prefix:
                    continue
                groups.setdefault(prefix, {})[suf] = col
                used.add(col)
                break
    groups = {k: v for k, v in groups.items() if len(v) >= 2}
    if len(groups) < 4:
        return None
    id_cols = [c for c in df.columns if c not in used]
    if not id_cols:
        return None
    id_col = id_cols[0]
    rows = []
    for _, rec in df.iterrows():
        for station, mapping in groups.items():
            item = {"批次": rec.get(id_col), "工位": station}
            for suf in UNPIVOT_SUFFIXES:
                if suf in mapping:
                    val = rec.get(mapping[suf])
                    item[suf] = pd.to_numeric(val, errors="coerce") if val not in (None, "－", "-") else None
            if all(item.get(s) in (None, 0) or (isinstance(item.get(s), float) and pd.isna(item.get(s))) for s in mapping):
                continue
            rows.append(item)
    if len(rows) < 8:
        return None
    long_df = pd.DataFrame(rows)
    columns = []
    for i, key in enumerate(long_df.columns):
        series = long_df[key]
        if key == "工位" or key == "批次":
            typ = "string"
        else:
            typ = "number"
            long_df[key] = pd.to_numeric(series, errors="coerce")
        nonempty = series.dropna()
        columns.append(
            {
                "key": key,
                "label": key,
                "index": i + 1,
                "type": typ,
                "fill": round(float(nonempty.shape[0] / max(len(long_df), 1)), 3),
                "unique": int(nonempty.nunique()) if typ != "number" else int(nonempty.nunique()),
                "samples": [ _sample(v) for v in nonempty.unique()[:6] ],
                "formula_cells": 0,
            }
        )
    preview = _json_records(long_df.head(8))
    return {
        "name": sheet["name"].strip() + "·展开",
        "hidden": False,
        "kind": "fact",
        "header_row": 1,
        "row_count": int(len(long_df)),
        "col_count": len(columns),
        "columns": columns,
        "formula": {"count": 0, "functions": [], "cross_sheets": [], "samples": [], "columns": []},
        "preview": preview,
        "df": long_df,
        "empty": False,
        "note": f"由「{sheet['name'].strip()}」的工位列自动展开（抽检数 / 不良数 / 良品数）。",
        "derived": True,
    }
